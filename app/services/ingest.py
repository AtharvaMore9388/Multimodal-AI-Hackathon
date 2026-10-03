import hashlib
import io
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core import storage
from app.services import topic_extractor, image_ocr

CHUNK_WORDS = 280
CHUNK_OVERLAP = 40

TEXT_EXTS = {".txt", ".md", ".markdown"}
DOC_EXTS = {".pdf", ".pptx", ".docx", ".doc"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".avi"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".aac"}


def _ext(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def _doc_id(filename: str, subject: str) -> str:
    return hashlib.sha1(f"{filename}:{subject}:{datetime.utcnow().isoformat()}".encode()).hexdigest()[:20]


def _chunk_id(doc_id: str, idx: int) -> str:
    return hashlib.sha1(f"{doc_id}:{idx}".encode()).hexdigest()[:20]


def _chunk_text(text: str, words_per_chunk: int = CHUNK_WORDS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    tokens = text.split()
    if len(tokens) <= words_per_chunk:
        return [text]
    chunks = []
    start = 0
    while start < len(tokens):
        end = min(start + words_per_chunk, len(tokens))
        chunks.append(" ".join(tokens[start:end]))
        if end == len(tokens):
            break
        start = end - overlap
    return chunks or [text]


def _read_plaintext(raw: bytes, filename: str) -> list[tuple[str, dict[str, Any]]]:
    ext = _ext(filename)
    if b"\x00" in raw[:2048]:
        return []
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = raw.decode("latin-1")
        except Exception:
            return []
    if text.count("\ufffd") > 5:
        return []
    chunks = _chunk_text(text)
    return [(c, {"page": None, "slide": None, "figures_count": 0, "figure_keywords": None}) for c in chunks]


def _read_pdf(raw: bytes) -> list[tuple[str, dict[str, Any]]]:
    out: list[tuple[str, dict[str, Any]]] = []
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""
                fig_info = image_ocr.scan_pdf_images(page, page_text)
                combined = page_text + ("\n" + fig_info["ocr_appended"] if fig_info["ocr_appended"] else "")
                if not combined.strip():
                    continue
                for ci, chunk in enumerate(_chunk_text(combined)):
                    meta = {
                        "page": page_no,
                        "slide": None,
                        "figures_count": fig_info["figures_count"] if ci == 0 else 0,
                        "figure_keywords": fig_info["figure_keywords"] if ci == 0 else None,
                    }
                    out.append((chunk, meta))
    except Exception:
        import pypdf
        try:
            reader = pypdf.PdfReader(io.BytesIO(raw))
            for i, p in enumerate(reader.pages, start=1):
                t = p.extract_text() or ""
                for ci, chunk in enumerate(_chunk_text(t)):
                    out.append((chunk, {"page": i, "slide": None, "figures_count": 0, "figure_keywords": None}))
        except Exception:
            return []
    return out


def _read_pptx(raw: bytes) -> list[tuple[str, dict[str, Any]]]:
    out: list[tuple[str, dict[str, Any]]] = []
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw))
        for slide_no, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for p in shape.text_frame.paragraphs:
                        parts.append(p.text)
                try:
                    fig_info = image_ocr.scan_pptx_picture(shape, " ".join(parts))
                    if fig_info["ocr_appended"]:
                        parts.append(fig_info["ocr_appended"])
                except Exception:
                    pass
            slide_text = "\n".join(parts)
            if not slide_text.strip():
                continue
            for ci, chunk in enumerate(_chunk_text(slide_text)):
                out.append((chunk, {"page": None, "slide": slide_no, "figures_count": 0, "figure_keywords": None}))
    except Exception:
        import zipfile
        import xml.etree.ElementTree as ET
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")])
                for slide_no, sf in enumerate(slide_files, start=1):
                    xml_content = z.read(sf)
                    root = ET.fromstring(xml_content)
                    texts = [elem.text for elem in root.iter() if elem.text and elem.text.strip()]
                    clean_text = " ".join(texts)
                    if clean_text.strip():
                        for ci, chunk in enumerate(_chunk_text(clean_text)):
                            out.append((chunk, {"page": None, "slide": slide_no, "figures_count": 0, "figure_keywords": None}))
        except Exception:
            pass
    return out


def _read_docx(raw: bytes) -> list[tuple[str, dict[str, Any]]]:
    out: list[tuple[str, dict[str, Any]]] = []
    try:
        from docx import Document
        doc = Document(io.BytesIO(raw))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paras)
        for ci, chunk in enumerate(_chunk_text(text)):
            out.append((chunk, {"page": ci + 1, "slide": None, "figures_count": 0, "figure_keywords": None}))
    except Exception:
        return _read_plaintext(raw, "fallback.txt")
    return out or _read_plaintext(raw, "fallback.txt")


def _read_media(raw: bytes, filename: str) -> list[tuple[str, dict[str, Any]]]:
    ext = _ext(filename)
    is_video = ext in VIDEO_EXTS
    label = "video" if is_video else "audio"
    transcript = _transcribe_media(raw, filename)
    if not transcript:
        transcript = f"Synthetic {label} transcript for {filename}. Upload a real {label} file to enable transcription-based retrieval. This placeholder describes the {label} as part of course lectures on {_filename_to_subject(filename)}."
    out: list[tuple[str, dict[str, Any]]] = []
    segments = _chunk_text(transcript, words_per_chunk=200, overlap=30)
    approx_dur_s = max(30, len(segments) * 25)
    for i, seg in enumerate(segments):
        start_s = i * 25
        end_s = min(approx_dur_s, (i + 1) * 25)
        out.append((seg, {
            "page": None, "slide": None,
            "video_start_s": float(start_s), "video_end_s": float(end_s),
            "figures_count": 0, "figure_keywords": None,
        }))
    return out


def _transcribe_media(raw: bytes, filename: str) -> str:
    try:
        import whisper
        model = whisper.load_model("tiny")
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=_ext(filename), delete=False) as f:
            f.write(raw)
            path = f.name
        try:
            res = model.transcribe(path, fp16=False)
            return res.get("text", "") or ""
        finally:
            try:
                os.unlink(path)
            except Exception:
                pass
    except Exception:
        return ""


def _filename_to_subject(filename: str) -> str:
    name = (filename or "").lower()
    if any(k in name for k in ["os", "operating", "paging", "deadlock", "banker"]):
        return "Operating Systems"
    if any(k in name for k in ["network", "tcp", "http", "dns"]):
        return "Computer Networks"
    if any(k in name for k in ["db", "database", "sql", "query", "acid"]):
        return "Database Systems"
    return "General"


def ingest_file(filename: str, raw: bytes, subject: str | None = None,
                user_email: str | None = None) -> dict[str, Any]:
    ext = _ext(filename)
    subject = subject or _filename_to_subject(filename)
    doc_id = _doc_id(filename, subject)
    uploaded_at = datetime.utcnow().isoformat()

    if ext in TEXT_EXTS:
        parsed = _read_plaintext(raw, filename)
    elif ext == ".pdf":
        parsed = _read_pdf(raw)
    elif ext in {".pptx", ".ppt"}:
        parsed = _read_pptx(raw)
    elif ext in {".docx", ".doc"}:
        parsed = _read_docx(raw)
    elif ext in VIDEO_EXTS or ext in AUDIO_EXTS:
        parsed = _read_media(raw, filename)
    else:
        parsed = _read_plaintext(raw, filename + ".txt")

    total_pages = max((m.get("page") or 0) for _, m in parsed) if parsed else 0
    total_slides = max((m.get("slide") or 0) for _, m in parsed) if parsed else 0
    duration_s = max((m.get("video_end_s") or 0) for _, m in parsed) if parsed else 0.0

    storage.add_document({
        "doc_id": doc_id, "filename": filename,
        "file_type": ext.lstrip("."), "subject": subject,
        "uploaded_at": uploaded_at, "user_email": user_email,
        "total_pages": total_pages, "total_slides": total_slides,
        "duration_seconds": duration_s,
    })

    topics_seen: set[str] = set()
    concepts_seen: set[str] = set()
    all_topics: list[dict[str, Any]] = []
    all_concepts: list[dict[str, Any]] = []
    total_chunks = len(parsed)

    for ci, (text, meta) in enumerate(parsed):
        chunk_id = _chunk_id(doc_id, ci)
        extracted = topic_extractor.extract_topics_and_concepts(text, subject, chunk_index=ci)
        chunk_topics = extracted["topics"]
        chunk_concepts = extracted["concepts"]
        for t in chunk_topics:
            storage.upsert_topic(t)
            if t["topic_id"] not in topics_seen:
                topics_seen.add(t["topic_id"])
                all_topics.append(t)
        for c in chunk_concepts:
            storage.add_concept(c)
            cid = f"{c['topic_id']}:{c['name']}"
            if cid not in concepts_seen:
                concepts_seen.add(cid)
                all_concepts.append(c)
        storage.add_doc_chunk({
            "chunk_id": chunk_id, "doc_id": doc_id, "chunk_index": ci, "text": text,
            "page": meta.get("page"), "slide": meta.get("slide"),
            "video_start_s": meta.get("video_start_s"), "video_end_s": meta.get("video_end_s"),
            "figures_count": meta.get("figures_count", 0),
            "figure_keywords": meta.get("figure_keywords"),
        })
        links = [(t["topic_id"], float(t.get("confidence", 0.7))) for t in chunk_topics]
        if links:
            storage.link_doc_chunk_to_topics(chunk_id, links)
        _store_embedding(chunk_id, text)

    return {
        "doc_id": doc_id,
        "filename": filename,
        "file_type": ext.lstrip("."),
        "subject": subject,
        "uploaded_at": uploaded_at,
        "chunks_ingested": total_chunks,
        "total_pages": total_pages,
        "total_slides": total_slides,
        "duration_seconds": duration_s,
        "topics_detected": all_topics,
        "concepts_detected": all_concepts,
    }


def _store_embedding(chunk_id: str, text: str) -> None:
    try:
        from sentence_transformers import SentenceTransformer
        import numpy as np
        model = _get_embed_model()
        if model is None:
            return
        vec = model.encode(text[:2000], convert_to_numpy=True, show_progress_bar=False)
        try:
            import chromadb
            client = chromadb.PersistentClient(path=str(Path(__file__).resolve().parent.parent / "chroma_data"))
            coll = client.get_or_create_collection("doc_chunks")
            coll.upsert(
                ids=[chunk_id],
                embeddings=[vec.tolist()],
                documents=[text[:1500]],
                metadatas=[{"chunk_id": chunk_id}],
            )
        except Exception:
            pass
    except Exception:
        pass


_EMBED_MODEL: Any = None


def _get_embed_model() -> Any:
    global _EMBED_MODEL
    if _EMBED_MODEL is not None:
        return _EMBED_MODEL
    try:
        from sentence_transformers import SentenceTransformer
        model_name = os.environ.get("EMBEDDING_MODEL", "all-MiniLM-L6-v2")
        _EMBED_MODEL = SentenceTransformer(model_name)
    except Exception:
        _EMBED_MODEL = False
    return _EMBED_MODEL if _EMBED_MODEL else None


def search_similar(query: str, top_k: int = 5) -> list[dict[str, Any]]:
    from app.core import storage
    rows: list[dict[str, Any]] = []
    conn = storage._connect()
    try:
        try:
            cur = conn.execute(
                "SELECT c.chunk_id, c.doc_id, c.text, c.page, c.slide, c.video_start_s, c.video_end_s, "
                "c.figures_count, c.figure_keywords, d.filename, d.file_type, d.subject "
                "FROM doc_chunks c LEFT JOIN documents d ON d.doc_id=c.doc_id ORDER BY c.rowid DESC LIMIT ?",
                (top_k * 10,),
            )
            rows = [dict(r) for r in cur.fetchall()]
        except Exception:
            pass
        if not rows:
            try:
                cur = conn.execute(
                    "SELECT chunk_id, doc_id, text, page, slide, video_start_s, video_end_s, "
                    "figures_count, figure_keywords FROM doc_chunks ORDER BY rowid DESC LIMIT ?",
                    (top_k * 10,),
                )
                for r in cur.fetchall():
                    dr = dict(r)
                    dr.setdefault("filename", "course_material.txt")
                    dr.setdefault("file_type", "txt")
                    dr.setdefault("subject", "General")
                    rows.append(dr)
            except Exception:
                pass
    finally:
        conn.close()

    if not rows:
        return []

    results: list[dict[str, Any]] = []
    try:
        model = _get_embed_model()
        if model and rows:
            import numpy as np
            q_vec = model.encode(query[:1000], convert_to_numpy=True, show_progress_bar=False)
            scored = []
            for r in rows:
                try:
                    d_vec = model.encode(r["text"][:2000], convert_to_numpy=True, show_progress_bar=False)
                    sim = float(np.dot(q_vec, d_vec) / (np.linalg.norm(q_vec) * np.linalg.norm(d_vec) + 1e-9))
                except Exception:
                    sim = 0.55
                scored.append((sim, r))
            scored.sort(key=lambda x: x[0], reverse=True)
            for sim, r in scored[:top_k]:
                r["relevance"] = max(0.25, min(1.0, 0.5 + sim / 2.0))
                r["source"] = r.get("filename") or "doc_chunks"
                results.append(r)
    except Exception:
        pass

    if not results:
        q_tokens = set(re.findall(r"[A-Za-z0-9_]+", query.lower()))
        scored = []
        for r in rows:
            t_tokens = set(re.findall(r"[A-Za-z0-9_]+", (r["text"] or "").lower()))
            overlap = len(q_tokens & t_tokens) / max(1, len(q_tokens))
            scored.append((overlap, r))
        scored.sort(key=lambda x: x[0], reverse=True)
        for overlap, r in scored[:top_k]:
            r["relevance"] = max(0.30, min(1.0, 0.45 + overlap * 0.45))
            r["source"] = r.get("filename") or "doc_chunks"
            results.append(r)
    return results
